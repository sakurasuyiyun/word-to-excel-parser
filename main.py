from docx import Document
import openpyxl
import re
import sys
from datetime import datetime
from pathlib import Path
from plugin.number import add_question_numbers_docx

def parse_word_to_excel(word_file, excel_file):
    doc = Document(word_file)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Exam Questions"
    sheet.append(["题目", "选项A", "选项B", "选项C", "选项D", "正确答案"])

    question_pattern = re.compile(r"^\d+\.\s*(.+)")
    answer_pattern = re.compile(r"^答案[:：]\s*(.+)")
    current_question = {}
    
    for paragraph in doc.paragraphs:
        line = paragraph.text.strip()
        if not line:
            continue

        question_match = question_pattern.match(line)
        if question_match:
            if current_question:
                sheet.append([
                    current_question.get("题目", ""),
                    current_question.get("选项A", ""),
                    current_question.get("选项B", ""),
                    current_question.get("选项C", ""),
                    current_question.get("选项D", ""),
                    current_question.get("答案", ""),
                ])
            current_question = {"题目": question_match.group(1).strip()}
            continue

        if line.startswith("A."):
            current_question["选项A"] = line[2:].strip()
        elif line.startswith("B."):
            current_question["选项B"] = line[2:].strip()
        elif line.startswith("C."):
            current_question["选项C"] = line[2:].strip()
        elif line.startswith("D."):
            current_question["选项D"] = line[2:].strip()

        answer_match = answer_pattern.match(line)
        if answer_match:  
            current_question["答案"] = answer_match.group(1).strip()

    
    if current_question:
        sheet.append([
            current_question.get("题目", ""),
            current_question.get("选项A", ""),
            current_question.get("选项B", ""),
            current_question.get("选项C", ""),
            current_question.get("选项D", ""),
            current_question.get("答案", ""),
        ])

    workbook.save(excel_file)
    print(f"成功将 Word 文件内容保存到 {excel_file}")
    
    

def check_number_isvalid(max_num, user_input):
    # 检查输入是否为非负整数
    if user_input.isdigit() == False:
        return False
    user_input = int(user_input)
    return 1 <= user_input <= max_num

def del_file(word_file):
    file_path = Path(word_file)
    if file_path.exists():  # 确保文件存在
        file_path.unlink()
    else:
        print(f"文件 {file_path} 不存在")


if __name__ == "__main__":
    # 根据当前时间戳生成文件名
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    input_path = "input.docx"
    output_docx_path = f"exam_{timestamp}.docx"
    output_excel_path = f"exam_{timestamp}.xlsx"

    print("请选择题型：1-选择题 2-判断题")
    question_type = input("请输入题型编号：").strip()

    check = check_number_isvalid(2, question_type)
    if(check == False):
        print("输入有误")
        sys.exit()

    print("是否需要添加题目序号？")
    add_nums = input("y-需要 n-不需要：").strip()
    if(add_nums == 'y'):
        add_question_numbers_docx(input_path, output_docx_path)
        parse_word_to_excel(output_docx_path, output_excel_path)
        del_file(output_docx_path)
    else:
        parse_word_to_excel(input_path, output_excel_path)
        