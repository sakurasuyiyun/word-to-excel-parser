from docx import Document
import openpyxl
import re

def multiple_choice_question(word_file, excel_file):
    doc = Document(word_file)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Exam Questions"
    sheet.append(["题目", "选项A", "选项B", "选项C", "选项D", "正确答案"])

    question_pattern = re.compile(r"^\d+\.\s*(.+)")
    answer_pattern = re.compile(r"^答案[:：]\s*(.+)")
    current_question = {}
    
    for paragraph in doc.paragraphs:
        line = paragraph.text.strip()
        if not line:
            continue

        question_match = question_pattern.match(line)
        if question_match:
            if current_question:
                sheet.append([
                    current_question.get("题目", ""),
                    current_question.get("选项A", ""),
                    current_question.get("选项B", ""),
                    current_question.get("选项C", ""),
                    current_question.get("选项D", ""),
                    current_question.get("答案", ""),
                ])
            current_question = {"题目": question_match.group(1).strip()}
            continue

        if line.startswith("A."):
            current_question["选项A"] = line[2:].strip()
        elif line.startswith("B."):
            current_question["选项B"] = line[2:].strip()
        elif line.startswith("C."):
            current_question["选项C"] = line[2:].strip()
        elif line.startswith("D."):
            current_question["选项D"] = line[2:].strip()

        answer_match = answer_pattern.match(line)
        if answer_match:  
            current_question["答案"] = answer_match.group(1).strip()

    
    if current_question:
        sheet.append([
            current_question.get("题目", ""),
            current_question.get("选项A", ""),
            current_question.get("选项B", ""),
            current_question.get("选项C", ""),
            current_question.get("选项D", ""),
            current_question.get("答案", ""),
        ])

    workbook.save(excel_file)
    print(f"成功将 Word 文件内容保存到 {excel_file}")
    